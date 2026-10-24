"""Check lab-report metadata underlines and native image/table captions. Render separately."""
import argparse,json,re
from docx import Document
from docx.oxml.ns import qn

def text_of(element):
    return ''.join(element.xpath('.//w:t/text()'))

def audit(path):
    d=Document(path);issues=[]
    for label in ['实验名称','班级','姓名','学号','座位号']:
        matches=[p for p in d.paragraphs[:10] if label in p.text]
        if not matches:
            issues.append(f'缺少信息栏 {label}');continue
        p=matches[0];start=p.text.index(label)+len(label)
        stop=len(p.text)
        for other in ['班级','姓名','学号','座位号']:
            pos=p.text.find(other,start)
            if pos>=0:stop=min(stop,pos)
        offset=0;value_runs=[];blank_underlined=False
        for r in p.runs:
            end=offset+len(r.text)
            segment=r.text[max(0,start-offset):max(0,stop-offset)] if end>start and offset<stop else ''
            if segment and not segment.strip(' ：:') and r.underline:blank_underlined=True
            if segment.strip(' ：:'):
                value_runs.append(r)
            offset=end
        if (not value_runs and not blank_underlined) or any(not r.underline for r in value_runs):
            issues.append(f'{label} 填写内容未全部使用真实下划线')
    elements=list(d._element.body);counts={'图':0,'表':0}
    for i,e in enumerate(elements):
        kind='表' if e.tag==qn('w:tbl') else ('图' if e.xpath('.//a:blip') else None)
        if kind is None:continue
        counts[kind]+=1;next_text=''
        for n in elements[i+1:]:
            if n.tag==qn('w:sectPr'):break
            t=text_of(n).strip()
            if t:next_text=t;break
            if n.tag==qn('w:tbl') or n.xpath('.//a:blip'):break
        if not re.match(rf'^{kind}\s*{counts[kind]}\s+\S',next_text):
            issues.append(f'{kind} {counts[kind]} 下方缺少按序编号及名称的题注')
    for i,t in enumerate(d.tables,1):
        width=sum(c.width or 0 for c in t.columns)
        available=max(s.page_width-s.left_margin-s.right_margin for s in d.sections)
        if width>available+9144:issues.append(f'表 {i} 的列宽超过最大版心')
    return {'file':str(path),'figures':counts['图'],'tables':counts['表'],'issues':issues,'scope':'只检查结构；仍需逐页渲染核对字体、图文布局和分页'}

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('document');a=p.parse_args()
    result=audit(a.document);print(json.dumps(result,ensure_ascii=False,indent=2));raise SystemExit(bool(result['issues']))
