"""Create the verified personal lab-report scaffold, without experimental content."""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def format_run(run, size=14, font='宋体', underline=False):
    run.font.name=font;run.font.size=Pt(size);run.bold=True;run.underline=underline
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)

def paragraph(doc,text='',size=14,center=False,heading=False):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1.5 if heading else 1.25
    p.paragraph_format.keep_with_next=heading
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    format_run(p.add_run(text),size)
    return p

def field(p,label,value='',padding=1):
    format_run(p.add_run(label),12)
    # Underline both the actual value and its adjacent spaces.
    format_run(p.add_run(' '*padding+value+' '*padding if value else '\u00a0'*max(3,padding*2)),12,underline=True)

def add_figure(doc,image,caption,width_cm=13.8):
    p=paragraph(doc,'',center=True);p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1
    p.add_run().add_picture(str(image),width=Cm(width_cm))
    return paragraph(doc,caption,12,True)

def add_wide_table(doc,headers,rows,caption,font_size=14,width_cm=None):
    """Headers/rows are already transposed by the caller when appropriate."""
    width_cm=width_cm or (doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)/360000
    count=len(headers);t=doc.add_table(rows=1,cols=count);t.autofit=False
    for c in t.columns:c.width=Cm(width_cm/count)
    for c,value in zip(t.rows[0].cells,headers):c.text=str(value)
    for row in rows:
        if len(row)!=count:raise ValueError('Table rows must match header count')
        for c,value in zip(t.add_row().cells,row):c.text=str(value)
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            c.width=Cm(width_cm/count);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');borders.append(e)
            props.append(borders)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(2)
                if ri==len(t.rows)-1:p.paragraph_format.keep_with_next=True
                for r in p.runs:format_run(r,font_size)
    paragraph(doc,caption,12,True)
    return t

def create_report(reference=None,experiment='',class_name='',name='',student_id='',seat='',sections=None):
    doc=Document(reference) if reference else Document()
    if not reference:
        for section in doc.sections:
            section.page_width=Cm(21);section.page_height=Cm(29.7)
            section.top_margin=Cm(2.54);section.bottom_margin=Cm(2.54)
            section.left_margin=Cm(3.175);section.right_margin=Cm(3.175)
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    for element in list(doc._element.body):
        if element.tag!=qn('w:sectPr'):doc._element.body.remove(element)
    p=paragraph(doc,'物理实验报告',15,True)
    format_run(p.runs[0],15,'黑体')
    p=paragraph(doc,'',12,True);field(p,'实验名称：',experiment,3)
    p=paragraph(doc,'',12)
    for label,value in [('班级：',class_name),('姓名：',name),('学号：',student_id),('座位号：',seat)]:
        field(p,label,value);format_run(p.add_run('  '),12)
    for title in sections or ['一、原始数据','二、数据分析','三、课后题']:
        paragraph(doc,title,16,heading=True)
    return doc

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',required=True);p.add_argument('--experiment',default='')
    p.add_argument('--seat',default='');p.add_argument('--name',default='');p.add_argument('--student-id',default='');p.add_argument('--class-name',default='');p.add_argument('--reference',default=str(Path(__file__).resolve().parents[1]/'assets/reference.docx'))
    a=p.parse_args();out=Path(a.output);out.parent.mkdir(parents=True,exist_ok=True)
    create_report(a.reference,a.experiment,class_name=a.class_name,name=a.name,student_id=a.student_id,seat=a.seat).save(out)
    print(out)
